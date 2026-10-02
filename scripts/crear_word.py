"""Versión Word editable del reporte existente, con el mismo contenido."""
import ast
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
datos = json.loads((ROOT / "output/mediciones.json").read_text(encoding="utf-8"))
source = ast.parse((ROOT / "scripts/crear_reporte.py").read_text(encoding="utf-8"))
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin, section.bottom_margin = Pt(43), Pt(54)
section.left_margin = section.right_margin = Pt(48)
section.footer_distance = Pt(25)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Arial", Pt(10)
normal.paragraph_format.line_spacing = Pt(14)
normal.paragraph_format.space_after = Pt(7)
title = doc.styles["Title"]
title.font.name, title.font.size = "Arial", Pt(23)
title.font.bold, title.font.color.rgb = True, RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(15)
title.paragraph_format.line_spacing = Pt(27)
heading = doc.styles["Heading 1"]
heading.font.name, heading.font.size = "Arial", Pt(12)
heading.font.bold, heading.font.color.rgb = True, RGBColor.from_string("163B37")
heading.paragraph_format.space_before = Pt(15)
heading.paragraph_format.space_after = Pt(7)
heading.paragraph_format.line_spacing = Pt(16)
heading.paragraph_format.keep_with_next = True
for style in doc.styles:
    for border in list(style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)


def agregar(texto, estilo="Normal"):
    p = doc.add_paragraph(style=estilo)
    for fragmento in re.split(r"(<super>.*?</super>)", texto):
        if fragmento.startswith("<super>"):
            r = p.add_run(fragmento[7:-8]); r.font.superscript = True
        else:
            p.add_run(fragmento)
    if estilo == "Caption":
        p.paragraph_format.line_spacing = Pt(12.5)
        p.paragraph_format.space_after = Pt(8)
        for r in p.runs:
            r.font.name, r.font.size = "Arial", Pt(9)
            r.font.italic = False
            r.font.bold = False
            r.font.color.rgb = RGBColor.from_string("52656A")
    return p


def tabla():
    filas = [["Dificultad", "Carreras", "Promedio (s)", "Mín. / máx. (s)", "Intentos promedio"]]
    for r in datos["resultados"]:
        filas.append([str(r["dificultad"]), str(datos["entorno"]["repeticiones"]), f"{r['promedio_s']:.3f}", f"{r['min_s']:.3f} / {r['max_s']:.3f}", f"{r['promedio_intentos']:,.1f}"])
    t = doc.add_table(rows=0, cols=5)
    t.autofit = False
    widths = [65, 53, 90, 115, 132]
    for col, w in zip(t.columns, widths): col.width = Pt(w)
    for i, fila in enumerate(filas):
        row = t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0: trpr.append(OxmlElement("w:tblHeader"))
        for cell, value, width in zip(row.cells, fila, widths):
            cell.width = Pt(width)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "163B37" if i == 0 else ("EDF4F2" if i % 2 else "FFFFFF"))
            cell._tc.get_or_add_tcPr().append(shade)
            margins = OxmlElement("w:tcMar")
            for side, value_margin in [("top", 170), ("bottom", 170), ("left", 85), ("right", 65)]:
                item = OxmlElement("w:" + side); item.set(qn("w:w"), str(value_margin)); item.set(qn("w:type"), "dxa"); margins.append(item)
            cell._tc.get_or_add_tcPr().append(margins)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(11)
            r = p.add_run(value); r.font.size = Pt(9); r.bold = i == 0
            if i == 0: r.font.color.rgb = RGBColor(255,255,255)


for node in source.body:
    if not isinstance(node, ast.Expr) or not isinstance(node.value, ast.Call): continue
    call = node.value
    if isinstance(call.func, ast.Name) and call.func.id in ("p", "h"):
        texto = eval(compile(ast.Expression(call.args[0]), "<texto>", "eval"), {"datos": datos})
        agregar(texto, "Heading 1" if call.func.id == "h" else "Normal")
    elif isinstance(call.func, ast.Attribute) and isinstance(call.func.value, ast.Name) and call.func.value.id == "contenido" and call.func.attr == "append":
        item = call.args[0]
        if isinstance(item, ast.Name) and item.id == "tabla": tabla()
        elif isinstance(item, ast.Call) and isinstance(item.func, ast.Name):
            if item.func.id == "PageBreak": doc.add_page_break()
            elif item.func.id == "Paragraph":
                texto = ast.literal_eval(item.args[0])
                estilo = ast.literal_eval(item.args[1].slice)
                agregar(texto, "Title" if estilo == "TituloLocal" else "Caption")
            elif item.func.id == "Spacer":
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = Pt(5); p.add_run().font.size = Pt(1)

footer = section.footer.paragraphs[0]
footer.paragraph_format.space_after = Pt(0)
footer.paragraph_format.tab_stops.add_tab_stop(Pt(470))
r = footer.add_run("Préstamos de laboratorio | Proof of Work\t"); r.font.size = Pt(8)
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
doc.core_properties.title = "Simulador de Proof of Work"
doc.core_properties.author = "Proyecto de Fundamentos de Blockchain"
destino = ROOT / "output/reporte_pow.docx"
doc.save(destino)
print(destino)
